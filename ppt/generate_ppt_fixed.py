"""
担保交易系统PPT生成器 - 修复版
使用 python-pptx 库生成专业的PowerPoint演示文稿
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_guarantee_trading_ppt():
    """创建担保交易系统PPT"""
    
    # 创建演示文稿
    prs = Presentation()
    prs.slide_width = Inches(16)  # 16:9 宽屏
    prs.slide_height = Inches(9)
    
    # 定义颜色方案
    PRIMARY_COLOR = RGBColor(102, 126, 234)  # #667eea
    SECONDARY_COLOR = RGBColor(118, 75, 162)  # #764ba2
    ACCENT_COLOR = RGBColor(39, 174, 96)  # #27ae60
    WARNING_COLOR = RGBColor(231, 76, 60)  # #e74c3c
    TEXT_COLOR = RGBColor(44, 62, 80)  # #2c3e50
    LIGHT_BG = RGBColor(248, 249, 250)  # #f8f9fa
    WHITE = RGBColor(255, 255, 255)
    
    # ==================== 第1页：封面 ====================
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # 空白布局
    add_gradient_background(slide1, PRIMARY_COLOR, SECONDARY_COLOR)
    
    # 标题
    title_box = slide1.shapes.add_textbox(Inches(1), Inches(3), Inches(14), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.text = "专业B2B担保交易解决方案"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(60)
    title_para.font.bold = True
    title_para.font.color.rgb = WHITE
    title_para.alignment = PP_ALIGN.CENTER
    
    # 副标题
    subtitle_box = slide1.shapes.add_textbox(Inches(1), Inches(4.8), Inches(14), Inches(0.8))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "让每一笔交易都安全可控"
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(36)
    subtitle_para.font.color.rgb = WHITE
    subtitle_para.alignment = PP_ALIGN.CENTER
    
    # 公司信息
    info_box = slide1.shapes.add_textbox(Inches(1), Inches(7), Inches(14), Inches(0.5))
    info_frame = info_box.text_frame
    info_frame.text = "您的公司 | 2025年11月"
    info_para = info_frame.paragraphs[0]
    info_para.font.size = Pt(24)
    info_para.font.color.rgb = RGBColor(200, 200, 200)
    info_para.alignment = PP_ALIGN.CENTER
    
    # ==================== 第2页：B2B交易的三大痛点 ====================
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide2, PRIMARY_COLOR, SECONDARY_COLOR)
    
    # 标题
    add_slide_title(slide2, "B2B交易的三大痛点")
    
    # 三列痛点
    pain_points = [
        {
            "emoji": "😰",
            "title": "买家的担忧",
            "points": [
                "付款后收不到货",
                "货物质量无法保障",
                "大额交易风险高"
            ]
        },
        {
            "emoji": "😥",
            "title": "卖家的担忧",
            "points": [
                "发货后收不到钱",
                "买家恶意拒收",
                "账期不确定"
            ]
        },
        {
            "emoji": "😓",
            "title": "平台的挑战",
            "points": [
                "交易纠纷处理成本高",
                "资金流转缺乏透明度",
                "缺乏有效风控手段"
            ]
        }
    ]
    
    for i, pain in enumerate(pain_points):
        x = 0.5 + i * 5.2
        box = add_rounded_rectangle(slide2, Inches(x), Inches(1.8), Inches(4.8), Inches(4.2), 
                                     LIGHT_BG)
        
        # 标题
        title_text = f"{pain['emoji']} {pain['title']}"
        tf = box.text_frame
        tf.text = title_text
        p = tf.paragraphs[0]
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = TEXT_COLOR
        
        # 痛点列表
        for point in pain['points']:
            p = tf.add_paragraph()
            p.text = f"├─ {point}"
            p.font.size = Pt(18)
            p.font.color.rgb = RGBColor(85, 85, 85)
            p.space_before = Pt(10)
            p.level = 0
    
    # 底部提示
    insight_box = add_rounded_rectangle(slide2, Inches(0.5), Inches(6.5), Inches(15), Inches(1), 
                                        RGBColor(255, 243, 205))
    tf = insight_box.text_frame
    tf.text = "💡 传统线下对公转账无法解决信任问题"
    p = tf.paragraphs[0]
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(133, 100, 4)
    p.alignment = PP_ALIGN.CENTER
    
    # ==================== 第3页：担保交易的商业价值 ====================
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide3, PRIMARY_COLOR, SECONDARY_COLOR)
    add_slide_title(slide3, "担保交易的商业价值")
    
    # 左右两个价值框
    values = [
        {
            "title": "对买卖双方的价值",
            "points": [
                "✓ 建立交易信任",
                "✓ 降低交易风险",
                "✓ 优化资金周转",
                "✓ 保障货款安全"
            ]
        },
        {
            "title": "对平台的价值",
            "points": [
                "✓ 提升平台GMV 35%+",
                "✓ 减少纠纷处理工作量 60%",
                "✓ 增加服务费收入（1-2%费率）",
                "✓ 提升客户粘性和复购率"
            ]
        }
    ]
    
    for i, value in enumerate(values):
        x = 0.5 + i * 7.5
        box = add_rounded_rectangle(slide3, Inches(x), Inches(1.8), Inches(7), Inches(3.5), 
                                     PRIMARY_COLOR)
        
        tf = box.text_frame
        tf.text = value['title']
        p = tf.paragraphs[0]
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = WHITE
        
        for point in value['points']:
            p = tf.add_paragraph()
            p.text = point
            p.font.size = Pt(20)
            p.font.color.rgb = WHITE
            p.space_before = Pt(12)
    
    # 底部数据亮点
    data_box = add_rounded_rectangle(slide3, Inches(0.5), Inches(5.8), Inches(15), Inches(1.2), 
                                     RGBColor(212, 237, 218))
    tf = data_box.text_frame
    tf.text = "📊 行业数据"
    p = tf.paragraphs[0]
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(21, 87, 36)
    p.alignment = PP_ALIGN.CENTER
    
    p = tf.add_paragraph()
    p.text = "使用担保交易的B2B平台，年交易额增长率达到45%"
    p.font.size = Pt(20)
    p.font.color.rgb = RGBColor(21, 87, 36)
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(8)
    
    # ==================== 第4页：担保交易流程对比 ====================
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide4, PRIMARY_COLOR, SECONDARY_COLOR)
    add_slide_title(slide4, "担保交易流程对比")
    
    # 传统流程（左侧）
    add_process_flow(slide4, 1, "传统流程（风险高）", WARNING_COLOR, [
        "买家付款 → 卖家收钱",
        "卖家发货（或不发）",
        "买家收货（维权困难）"
    ])
    
    # 担保流程（右侧）
    add_process_flow(slide4, 8.5, "我们的担保流程（安全可控）", ACCENT_COLOR, [
        "买家付款 → 平台托管",
        "卖家发货",
        "买家验收 → 平台释放资金",
        "交易完成（全程可追溯）"
    ])
    
    # 底部控制点
    control_box = add_rounded_rectangle(slide4, Inches(0.5), Inches(6.5), Inches(15), Inches(0.8), 
                                        RGBColor(255, 243, 205))
    tf = control_box.text_frame
    tf.text = "⏱️ 关键控制点：收款确认、验收仲裁、资金释放"
    p = tf.paragraphs[0]
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(133, 100, 4)
    p.alignment = PP_ALIGN.CENTER
    
    # ==================== 第5页：系统核心功能矩阵 ====================
    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide5, PRIMARY_COLOR, SECONDARY_COLOR)
    add_slide_title(slide5, "系统核心功能矩阵")
    
    # 创建表格
    rows = 5
    cols = 3
    left = Inches(0.5)
    top = Inches(1.8)
    width = Inches(15)
    height = Inches(5)
    
    table = slide5.shapes.add_table(rows, cols, left, top, width, height).table
    
    # 设置列宽
    table.columns[0].width = Inches(3)
    table.columns[1].width = Inches(4)
    table.columns[2].width = Inches(8)
    
    # 表头
    headers = ["功能模块", "核心能力", "行业领先特性"]
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY_COLOR
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
    
    # 表格内容
    data = [
        ["资金托管", "线下汇款确认机制", "• 水单识别与验证\n• 多银行账户管理\n• 资金池实时监控"],
        ["智能风控", "多维度风险评估", "• 订单风险评分\n• 黑名单预警\n• 异常交易拦截"],
        ["争议仲裁", "专业化纠纷处理", "• 证据链管理\n• 三方沟通机制\n• 仲裁结果可追溯"],
        ["数据看板", "实时运营监控", "• 资金流量分析\n• 效能指标追踪\n• 风险预警系统"]
    ]
    
    for i, row_data in enumerate(data):
        for j, cell_text in enumerate(row_data):
            cell = table.cell(i + 1, j)
            cell.text = cell_text
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(16)
            p.font.color.rgb = TEXT_COLOR
            if j == 0:
                p.font.bold = True
    
    # ==================== 第6页：六大核心优势 ====================
    slide6 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide6, PRIMARY_COLOR, SECONDARY_COLOR)
    add_slide_title(slide6, "六大核心优势")
    
    advantages = [
        {"title": "1️⃣ 本地化适配", "points": ["支持线下对公转账", "符合国内财务合规", "多币种结算支持"]},
        {"title": "2️⃣ 灵活的资金控制", "points": ["自定义账期管理", "分阶段资金释放", "灵活的手续费规则"]},
        {"title": "3️⃣ 智能化风控", "points": ["AI风险评分", "实时异常监测", "黑名单自动拦截"]},
        {"title": "4️⃣ 高效的争议处理", "points": ["平均处理周期<3天", "证据链完整保存", "多方协商机制"]},
        {"title": "5️⃣ 专业化报表", "points": ["实时数据看板", "自动化报告生成", "预警推送机制"]},
        {"title": "6️⃣ 极致的用户体验", "points": ["收款确认<2小时", "移动端随时监控", "一站式操作界面"]}
    ]
    
    for i in range(2):
        for j in range(3):
            idx = i * 3 + j
            if idx < len(advantages):
                x = 0.5 + j * 5.2
                y = 1.8 + i * 2.5
                
                box = add_rounded_rectangle(slide6, Inches(x), Inches(y), Inches(4.8), Inches(2.2), 
                                            RGBColor(240, 147, 251))
                
                adv = advantages[idx]
                tf = box.text_frame
                tf.text = adv['title']
                p = tf.paragraphs[0]
                p.font.size = Pt(22)
                p.font.bold = True
                p.font.color.rgb = WHITE
                
                for point in adv['points']:
                    p = tf.add_paragraph()
                    p.text = f"• {point}"
                    p.font.size = Pt(16)
                    p.font.color.rgb = WHITE
                    p.space_before = Pt(8)
    
    # ==================== 第7页：技术架构亮点 ====================
    slide7 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide7, PRIMARY_COLOR, SECONDARY_COLOR)
    add_slide_title(slide7, "技术架构亮点")
    
    arch_layers = [
        {"title": "安全层", "items": ["多因素认证(MFA)", "KYC企业认证", "数据加密传输", "反洗钱(AML)合规"]},
        {"title": "业务层", "items": ["订单状态机管理", "智能工作流引擎", "自动化规则引擎", "Webhook事件推送"]},
        {"title": "数据层", "items": ["实时数据分析", "资金流追踪", "审计日志系统", "BI报表引擎"]}
    ]
    
    for i, layer in enumerate(arch_layers):
        x = 0.5 + i * 5.2
        box = add_rounded_rectangle(slide7, Inches(x), Inches(1.8), Inches(4.8), Inches(3.5), 
                                     PRIMARY_COLOR)
        
        tf = box.text_frame
        tf.text = layer['title']
        p = tf.paragraphs[0]
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        
        for item in layer['items']:
            p = tf.add_paragraph()
            p.text = f"├─ {item}"
            p.font.size = Pt(16)
            p.font.color.rgb = WHITE
            p.space_before = Pt(10)
    
    # 底部统计
    stats_box = add_rounded_rectangle(slide7, Inches(0.5), Inches(5.8), Inches(15), Inches(0.8), 
                                      LIGHT_BG)
    tf = stats_box.text_frame
    tf.text = "🔒 银行级安全标准 | 99.9%系统可用性 | 支持10万+日订单处理"
    p = tf.paragraphs[0]
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = TEXT_COLOR
    p.alignment = PP_ALIGN.CENTER
    
    # ==================== 第8页：与知名系统的对标 ====================
    slide8 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide8, PRIMARY_COLOR, SECONDARY_COLOR)
    add_slide_title(slide8, "与知名系统的对标")
    
    # 创建对比表格
    rows = 8
    cols = 4
    table = slide8.shapes.add_table(rows, cols, Inches(0.5), Inches(1.8), Inches(15), Inches(5)).table
    
    # 设置列宽
    table.columns[0].width = Inches(3)
    for i in range(1, 4):
        table.columns[i].width = Inches(4)
    
    # 表头
    comparison_headers = ["对比维度", "Escrow.com", "支付宝担保", "我们的方案"]
    for i, header in enumerate(comparison_headers):
        cell = table.cell(0, i)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = TEXT_COLOR
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
    
    # 对比数据
    comparison_data = [
        ["适用场景", "国际大额交易", "C2C小额", "B2B中大额 ✓"],
        ["本地化", "弱", "强", "强 ✓"],
        ["线下汇款支持", "❌", "❌", "✓"],
        ["自定义风控", "有限", "无", "完全自定义 ✓"],
        ["数据看板", "基础", "无", "专业级 ✓"],
        ["部署方式", "SaaS", "托管", "私有部署/SaaS ✓"],
        ["服务费", "0.89-5%", "免费", "1-2%（可协商）"]
    ]
    
    for i, row_data in enumerate(comparison_data):
        for j, cell_text in enumerate(row_data):
            cell = table.cell(i + 1, j)
            cell.text = cell_text
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(16)
            p.alignment = PP_ALIGN.CENTER
            
            if j == 0:
                p.font.bold = True
                p.font.color.rgb = TEXT_COLOR
            elif j == 3:  # 我们的方案列高亮
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(212, 237, 218)
                p.font.bold = True
                p.font.color.rgb = RGBColor(21, 87, 36)
    
    # ==================== 第9页：实时运营监控看板 ====================
    slide9 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide9, PRIMARY_COLOR, SECONDARY_COLOR)
    add_slide_title(slide9, "实时运营监控看板")
    
    # 资金池总览
    fund_box = add_rounded_rectangle(slide9, Inches(0.5), Inches(1.8), Inches(15), Inches(2.8), 
                                     PRIMARY_COLOR)
    tf = fund_box.text_frame
    tf.text = "💰 今日资金池状态总览"
    p = tf.paragraphs[0]
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    p = tf.add_paragraph()
    p.text = "监管中总资金：¥8,456,000"
    p.font.size = Pt(42)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(20)
    
    # 资金分解（文字说明）
    p = tf.add_paragraph()
    p.text = "待确认收款 ¥1,234,000 (45笔)  |  已确认待发货 ¥2,100,000 (78笔)"
    p.font.size = Pt(18)
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(15)
    
    p = tf.add_paragraph()
    p.text = "已发货待验收 ¥3,890,000 (132笔)  |  验收通过待结算 ¥1,232,000 (56笔)"
    p.font.size = Pt(18)
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(8)
    
    # 底部关键指标
    metrics = [
        {"label": "平均资金占用周期", "value": "12.5天"},
        {"label": "今日资金净流入", "value": "+¥180,000"},
        {"label": "本周结算准时率", "value": "89.5%"}
    ]
    
    for i, metric in enumerate(metrics):
        x = 0.5 + i * 5.2
        metric_box = add_rounded_rectangle(slide9, Inches(x), Inches(5.2), Inches(4.8), Inches(1.2), 
                                           WHITE)
        
        tf = metric_box.text_frame
        tf.text = metric['label']
        p = tf.paragraphs[0]
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(127, 140, 141)
        p.alignment = PP_ALIGN.CENTER
        
        p = tf.add_paragraph()
        p.text = metric['value']
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = TEXT_COLOR
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(8)
    
    # 底部洞察
    insight_box = add_rounded_rectangle(slide9, Inches(0.5), Inches(6.7), Inches(15), Inches(0.6), 
                                        RGBColor(255, 243, 205))
    tf = insight_box.text_frame
    tf.text = "💡 管理者5秒内掌握全局资金状态"
    p = tf.paragraphs[0]
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(133, 100, 4)
    p.alignment = PP_ALIGN.CENTER
    
    # ==================== 第10页：结束页 ====================
    slide10 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide10, PRIMARY_COLOR, SECONDARY_COLOR)
    
    # 感谢文字
    thanks_box = slide10.shapes.add_textbox(Inches(1), Inches(3), Inches(14), Inches(1))
    tf = thanks_box.text_frame
    tf.text = "感谢观看"
    p = tf.paragraphs[0]
    p.font.size = Pt(64)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # 联系方式
    contact_info = [
        "📞 联系电话：400-XXX-XXXX",
        "📧 邮箱：contact@yourcompany.com",
        "🌐 官网：www.yourcompany.com"
    ]
    
    contact_box = slide10.shapes.add_textbox(Inches(1), Inches(4.5), Inches(14), Inches(2))
    tf = contact_box.text_frame
    
    for i, info in enumerate(contact_info):
        if i > 0:
            p = tf.add_paragraph()
        else:
            p = tf.paragraphs[0]
        p.text = info
        p.font.size = Pt(28)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(15)
    
    # 底部标语
    tagline_box = slide10.shapes.add_textbox(Inches(1), Inches(7), Inches(14), Inches(0.8))
    tf = tagline_box.text_frame
    tf.text = "让我们一起打造更安全的B2B交易生态"
    p = tf.paragraphs[0]
    p.font.size = Pt(24)
    p.font.color.rgb = RGBColor(200, 200, 200)
    p.alignment = PP_ALIGN.CENTER
    
    # 保存文件
    filename = "担保交易系统解决方案_专业版.pptx"
    prs.save(filename)
    print(f"✅ PPT生成成功！文件名：{filename}")
    print(f"📄 共生成 {len(prs.slides)} 页幻灯片")
    return filename


# ==================== 辅助函数 ====================

def add_gradient_background(slide, color1, color2):
    """添加渐变背景"""
    background = slide.background
    fill = background.fill
    fill.gradient()
    fill.gradient_angle = 135
    fill.gradient_stops[0].color.rgb = color1
    fill.gradient_stops[1].color.rgb = color2


def add_slide_title(slide, title_text):
    """添加幻灯片标题"""
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(15), Inches(0.8))
    tf = title_box.text_frame
    tf.text = title_text
    p = tf.paragraphs[0]
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)


def add_rounded_rectangle(slide, left, top, width, height, fill_color):
    """添加圆角矩形"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape


def add_process_flow(slide, x_start, title, border_color, steps):
    """添加流程图"""
    # 标题框
    title_box = add_rounded_rectangle(slide, Inches(x_start), Inches(1.8), Inches(6.5), Inches(0.6), 
                                      RGBColor(248, 249, 250))
    title_box.line.color.rgb = border_color
    title_box.line.width = Pt(3)
    
    tf = title_box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = border_color
    p.alignment = PP_ALIGN.CENTER
    
    # 流程步骤
    for i, step in enumerate(steps):
        y = 2.6 + i * 0.9
        
        step_box = add_rounded_rectangle(slide, Inches(x_start), Inches(y), Inches(6.5), Inches(0.7), 
                                         RGBColor(255, 255, 255))
        
        tf = step_box.text_frame
        tf.text = step
        p = tf.paragraphs[0]
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(44, 62, 80)
        p.alignment = PP_ALIGN.CENTER
        
        # 添加箭头（除了最后一个步骤）
        if i < len(steps) - 1:
            arrow_y = y + 0.75
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.DOWN_ARROW,
                Inches(x_start + 3), Inches(arrow_y),
                Inches(0.5), Inches(0.3)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor(149, 165, 166)
            arrow.line.fill.background()


# ==================== 主程序 ====================

if __name__ == "__main__":
    print("🚀 开始生成担保交易系统PPT...")
    print("📦 需要安装 python-pptx 库")
    print("   安装命令: pip install python-pptx")
    print("-" * 50)
    
    try:
        create_guarantee_trading_ppt()
        print("\n🎉 恭喜！PPT生成完成！")
        print("💡 提示：你可以用 PowerPoint 或 WPS 打开查看")
    except ImportError:
        print("\n❌ 错误：未安装 python-pptx 库")
        print("请先运行: pip install python-pptx")
    except Exception as e:
        print(f"\n❌ 生成失败：{e}")
        import traceback
        traceback.print_exc()